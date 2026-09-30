#!/usr/bin/env python3
"""
source_wave3.py - one-off AI Ark sourcing job for RolledCo outreach wave 3.

    ClaudeCode_ArkAI_API=... python source_wave3.py            # full run
    python source_wave3.py --per-segment 5                     # smoke test

Reuses the AIArk client (rate limit, retry, response cache) from enrich_contacts.py.
Per segment: run the people search (0.5 credit / hit), skip anyone already in an
outreach file, then export the person by id (1 credit when a verified email is
found, 0 otherwise). Keep only VALID, non-free (work) emails. Stop when the
segment has --per-segment people or the page cap is hit.

Output: rolledco_outreach_wave3.csv with the tracker columns
  Name, Title, Company, LinkedIn URL, Source Variant, Email, Email Confidence,
  Phone, Enrichment Status  + Segment + Hook
Nothing is emailed. Phone finder is not called.
"""
import argparse
import csv
import glob
import json
import os
import re
import sys
import unicodedata
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from enrich_contacts import AIArk, API_KEY, merge, clean_linkedin, pick_email, unwrap  # noqa: E402

OUT = HERE / "rolledco_outreach_wave3.csv"
CACHE = HERE / ".aiark_cache_wave3.json"
US = {"location": {"any": {"include": ["United States"]}}}
PAGE_SIZE = 25

TRACKER_FIELDS = ["Name", "Title", "Company", "LinkedIn URL", "Source Variant", "Email", "Email Confidence", "Phone", "Enrichment Status"]
FIELDS = TRACKER_FIELDS + ["Segment", "Hook"]


# ------------------------------------------------------------------ dedupe set

def fold_name(s):
    s = unicodedata.normalize("NFKD", s or "").encode("ascii", "ignore").decode().lower()
    s = re.sub(r"\(.*?\)", "", s)
    s = re.sub(r"\b(jr|sr|ii|iii|mba|cfa|cpa|phd)\b\.?", " ", s)
    s = re.sub(r'["“”]', " ", s)
    parts = [p for p in re.sub(r"[^a-z ]+", " ", s).split() if len(p) > 1]
    return " ".join(parts)


def load_already_contacted():
    """Every outreach file we have: the 4 EF Start input CSVs, the enriched tracker (same people
    plus emails/linkedin), and the two outreach PDFs exported from Drive (names only)."""
    names, links, emails, sources = {}, {}, {}, []
    rows, per_file = merge(str(HERE / "input"))
    for r in rows:
        names[fold_name(r["name"])] = "input CSVs"
        if r["linkedin_url"]:
            links[r["linkedin_url"].lower()] = "input CSVs"
    sources.append(f"enrich/input/*.csv: {len(rows)} people")
    for p in [HERE / "EF_Start_contacts_dry_run.xlsx", HERE / "contacts_enriched.xlsx"]:
        if not p.exists():
            continue
        from openpyxl import load_workbook
        ws = load_workbook(p, read_only=True)["Contacts"]
        n = 0
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0 or not row or not row[0]:
                continue
            n += 1
            names.setdefault(fold_name(str(row[0])), p.name)
            if row[3]:
                links.setdefault(str(row[3]).lower().rstrip("/"), p.name)
            if len(row) > 5 and row[5]:
                emails[str(row[5]).lower()] = p.name
        sources.append(f"{p.name}: {n} people")
    for p in [HERE / ".outreach_sent_pdf.txt", HERE / ".outreach_targets_pdf.txt"]:
        if not p.exists():
            continue
        t = p.read_text()
        found = [n.strip() for n in re.findall(r"\*\*([^*\n]{3,60})\*\*", t)]
        found = [n for n in found if len(n.split()) >= 2 and not re.search(r"Outreach|Page \d|operators|faculty|Holdco|Software", n)]
        if not found:  # "Sent" pdf has no bold markers: "NN First Last Title · Company"
            found = re.findall(r"(?:^|\n)(?:\d{2} )?([A-Z][\w'\-]+(?: [A-Z][\w'\-\.]+){1,3}) (?:Co-founder|Founder|CEO|President|COO|Owner|Partner|Lecturer|Adjunct|Co-CEO|Operator|Co-owner)", t)
        for n in found:
            names.setdefault(fold_name(n), p.name)
        sources.append(f"{p.name}: {len(set(found))} names")
    return names, links, emails, sources


# ------------------------------------------------------------------ segments

def title_any(terms, mode="WORD"):
    return {"experience": {"current": {"title": {"any": {"include": {"mode": mode, "content": terms}}}}}}


def kw(source, terms, mode="SMART"):
    return {"any": {"include": {"sources": [{"mode": mode, "source": source}], "content": terms}}}


SCHOOL_NAMES = ["Stanford Graduate School of Business", "Harvard Business School", "The Wharton School",
                "Kellogg School of Management", "Columbia Business School", "NYU Stern School of Business"]
LEADER_TITLES = ["Founder", "Co-Founder", "CEO", "Chief Executive Officer", "Operating Partner", "Managing Partner", "President", "Owner", "Principal"]
TRADES = ["HVAC", "plumbing", "pest control", "landscaping", "lawn care", "home services", "accounting firms", "CPA firms", "bookkeeping", "tax and accounting"]
SIZE = {"type": "RANGE", "range": [{"start": 2, "end": 1000}]}

EXCL_SBA = r"clos(ing|er)|processor|underwrit|analyst|assistant|coordinator|servic(ing|er)|documentation|compliance|operations|intern\b|packag|admin|paralegal|credit|marketing|recruit|talent|specialist|human resources|\bIT\b|engineer|product|customer experience|asset based|investment advisor|wine|beverage|middle market|emerging markets|counsel|legal|accountant|treasur|data"
EXCL_SEARCH = r"job search|executive search|title search|talent|recruit|staffing|placement|headhunt|search engine|\bSEO\b|paid search|research|marketing|estimated|search consultant|retained search|analyst|intern\b|patent|medical|library|writer|associate\b|assistant|coordinator|student|title company|first american|abstract"
EXCL_LEADER = r"regional|market area|market president|division|area president|broker|realtor|general manager|branch|vice president|\bvp\b|sales|marketing|assistant|associate|analyst|intern\b|payroll|\bhr\b|human resources|recruit|controller|specialist|coordinator|account executive|manager\b|product owner|counsel|legal|engineer|technician|homeowner|business owner\b.*(?:realtor|agent)"
LEADER_RE = r"\b(founder|co-founder|ceo|chief executive|operating partner|managing partner|owner|principal|chairman|president)\b"
TRADES_RE = r"home service|hvac|heating|air conditioning|plumb|pest|landscap|lawn|tree care|accounting|\bcpa\b|bookkeep|\btax\b|electrical|roofing|restoration|cleaning|garage door|irrigation|pool service"


def keep_sba(p):
    t = cur_title(p)
    return re.search(r"\bSBA\b|acquisition|lend|loan officer|business development|\bBDO\b|7\(a\)|M&A", t, re.I) and not re.search(EXCL_SBA, t, re.I)


def keep_search(p):
    t, h, ct = cur_title(p), headline(p), company_text(p)
    if re.search(EXCL_SEARCH, t + " " + h + " " + company_name(p), re.I):
        return False
    if re.search(r"executive search|recruit|staffing|placement|headhunt|talent|retained search|search firm|title insurance|title search", ct, re.I):
        return False
    eta = r"search fund|searcher|acquisition entrepreneur|entrepreneur(ship)? through acquisition|\bETA\b|self[- ]funded|search partner|search capital|acquisition partners"
    if re.search(eta, t + " " + h, re.I):
        return True
    return bool(re.search(r"\b(ceo|chief executive|president|founder)\b", t, re.I) and re.search(r"search fund|entrepreneurship through acquisition", ct, re.I))


def employees(p):
    fin = (p.get("company") or {}).get("financial") or {}
    n = fin.get("employees") if isinstance(fin, dict) else None
    if isinstance(n, dict):
        n = n.get("end") or n.get("start")
    for g in p.get("position_groups") or []:
        if (g.get("date") or {}).get("end") is None:
            e = (g.get("company") or {}).get("employees") or {}
            if isinstance(e, dict) and e.get("start"):
                n = n or e.get("start")
            break
    return n if isinstance(n, int) else 0


ACQ_RE = r"acqui|roll[- ]?up|platform|holding compan|holdco|consolidat|family of (?:brands|businesses|companies)|group of (?:home service|companies|businesses)|portfolio compan|partner(?:s|ing)? with (?:owners|founders|operators)|buy(?:s|ing)? (?:and|&) (?:grow|operate|build)|multi[- ]brand|operating compan|succession"
EXCL_CO = r"real estate|realty|broker|home warranty|\bapp\b|on-demand|marketing|digital agency|lead generation|health ?care|hospice|in-home care|home care|senior care|nursing|staffing|distribut|software|saas|franchise sales|insurance|mortgage|construction materials|manufactur|wholesale"


def keep_holdco(p):
    t = cur_title(p)
    if employees(p) > 1000:
        return False
    ct = company_text(p)
    if re.search(EXCL_CO, ct, re.I) and not re.search(r"acqui|roll[- ]?up|holding compan", ct, re.I):
        return False
    if re.search(r"real estate|realty|home warranty|health ?care|hospice|home care|senior care|software|saas|distribut|manufactur", ct, re.I):
        return False
    return bool(re.search(LEADER_RE, t, re.I) and not re.search(EXCL_LEADER, t, re.I)
                and re.search(TRADES_RE, ct, re.I) and re.search(ACQ_RE, ct, re.I))


SEGMENTS = [
    {
        "name": "SBA acquisition lenders",
        "queries": [
            {"label": "sba-live-oak", "max": 20,
             "body": {"contact": {**title_any(["SBA", "Business Development", "Lending", "Acquisition", "Loan Officer"]), **US},
                      "account": {"name": {"any": {"include": {"mode": "WORD", "content": ["Live Oak Bank"]}}}}}},
            {"label": "sba-titles",
             "body": {"contact": {**title_any(["SBA Business Development Officer", "SBA BDO", "SBA Loan Officer", "SBA Lending", "SBA Lender",
                                              "SBA Business Development", "Acquisition Lending", "SBA Relationship Manager", "SBA Acquisition"]), **US}}},
        ],
        "keep": keep_sba,
    },
    {
        "name": "Search funders and ETA buyers",
        "queries": [
            {"label": "searcher-title",
             "body": {"contact": {**title_any(["Search Fund", "Acquisition Entrepreneur", "Entrepreneur Through Acquisition", "Entrepreneurship Through Acquisition",
                                              "Search Fund Principal", "Search Fund Entrepreneur", "Self-Funded Searcher", "Searcher"]), **US}}},
            {"label": "searchfund-company",
             "body": {"contact": {**title_any(LEADER_TITLES + ["Searcher"]), **US},
                      "account": {"keyword": kw("NAME", ["search fund", "search capital", "acquisition partners", "succession partners", "legacy partners"], "WORD")}}},
            {"label": "searcher-headline",
             "body": {"contact": {"keyword": kw("HEADLINE", ["search fund", "self-funded searcher", "acquisition entrepreneur", "ETA searcher", "searching to acquire"], "WORD"), **US}}},
            {"label": "ceo-of-searchfund-acquired",
             "body": {"contact": {**title_any(["CEO", "President", "Chief Executive Officer"]), **US},
                      "account": {"keyword": kw("DESCRIPTION", ["search fund", "entrepreneurship through acquisition"], "WORD"), "employeeSize": SIZE}}},
        ],
        "keep": keep_search,
    },
    {
        "name": "Small holdco and roll-up operators",
        "queries": [
            {"label": "holdco-acquires-trades",
             "body": {"contact": {**title_any(LEADER_TITLES), **US},
                      "account": {"keyword": {"all": {"include": {"sources": [{"mode": "SMART", "source": "DESCRIPTION"}], "content": ["acquires businesses"]}},
                                              "any": {"include": {"sources": [{"mode": "WORD", "source": "DESCRIPTION"}], "content": TRADES}}},
                                  "employeeSize": SIZE}}},
            {"label": "holdco-rollup-trades",
             "body": {"contact": {**title_any(LEADER_TITLES), **US},
                      "account": {"keyword": {"all": {"include": {"sources": [{"mode": "SMART", "source": "DESCRIPTION"}], "content": ["roll-up platform"]}},
                                              "any": {"include": {"sources": [{"mode": "WORD", "source": "DESCRIPTION"}], "content": TRADES}}},
                                  "employeeSize": SIZE}}},
            {"label": "holdco-name-trades",
             "body": {"contact": {**title_any(LEADER_TITLES), **US},
                      "account": {"name": {"any": {"include": {"mode": "WORD", "content": ["Holdings", "Holdco", "Service Partners", "Home Services", "Services Group", "Service Group"]}}},
                                  "keyword": kw("DESCRIPTION", TRADES, "WORD"), "employeeSize": SIZE}}},
            {"label": "holdco-accounting",
             "body": {"contact": {**title_any(LEADER_TITLES), **US},
                      "account": {"keyword": {"all": {"include": {"sources": [{"mode": "SMART", "source": "DESCRIPTION"}], "content": ["acquiring accounting firms"]}}},
                                  "employeeSize": SIZE}}},
        ],
        "keep": keep_holdco,
    },
    {
        "name": "MBA ETA club leaders",
        "queries": [
            {"label": "eta-club-title",
             "body": {"contact": {**title_any(["ETA Club", "Entrepreneurship Through Acquisition Club", "Search Fund Club", "ETA & Search Fund Club",
                                              "Entrepreneurship Through Acquisition", "Search Fund and ETA Club", "ETA and Search Fund Club"]), **US}}},
            {"label": "eta-club-org",
             "body": {"contact": {"keyword": kw("ORGANIZATION", ["Entrepreneurship Through Acquisition", "ETA Club", "Search Fund Club"], "WORD"), **US}}},
            {"label": "eta-club-school",
             "body": {"contact": {**title_any(["Co-President", "President", "Vice President", "VP", "Co-Chair"]),
                                  "keyword": kw("HEADLINE", ["ETA", "Entrepreneurship Through Acquisition", "Search Fund"], "WORD"), **US},
                      "account": {"name": {"any": {"include": {"mode": "WORD", "content": SCHOOL_NAMES}}}}}},
        ],
        "keep": lambda p: club_hit(p) is not None,
    },
]

SCHOOLS = [
    ("Stanford GSB", r"stanford"),
    ("HBS", r"harvard business|\bhbs\b"),
    ("Wharton", r"wharton|university of pennsylvania"),
    ("Kellogg", r"kellogg|northwestern university"),
    ("Columbia Business School", r"columbia business|columbia university|\bcbs\b"),
    ("NYU Stern", r"nyu stern|\bnyu\b|new york university"),
]
CLUB_RE = re.compile(r"entrepreneurship through acquisition|\bETA\b|search fund", re.I)
LEAD_RE = re.compile(r"president|\bvp\b|vice president|chair|lead|director|head|founder|board", re.I)


def positions(p):
    out = []
    for g in p.get("position_groups") or []:
        comp = (g.get("company") or {}).get("name", "") or ""
        for pos in g.get("profile_positions") or []:
            out.append((pos.get("title") or "", comp or pos.get("company") or "", (pos.get("date") or {})))
    return out


def cur_title(p):
    return (p.get("profile") or {}).get("title") or ""


def headline(p):
    return (p.get("profile") or {}).get("headline") or ""


def company_name(p):
    c = p.get("company") or {}
    s = c.get("summary") or {}
    if isinstance(s, dict) and s.get("name"):
        return s["name"]
    for g in p.get("position_groups") or []:
        if (g.get("date") or {}).get("end") is None:
            return (g.get("company") or {}).get("name", "") or ""
    return ""


def company_text(p):
    c = p.get("company") or {}
    s = c.get("summary") or {}
    bits = [company_name(p), headline(p), s.get("description") or "" if isinstance(s, dict) else "", " ".join(c.get("keywords") or []) if isinstance(c.get("keywords"), list) else ""]
    return " ".join(str(b) for b in bits if b)


def person_text(p):
    return " ".join([cur_title(p), headline(p), (p.get("profile") or {}).get("summary") or "",
                     " ".join(f"{t} {c}" for t, c, _ in positions(p)),
                     " ".join(((o or {}).get("name") or "") + " " + ((o or {}).get("position") or "") for o in p.get("organizations") or []),
                     " ".join(((e.get("school") or {}).get("name") or "") for e in p.get("educations") or [])])


def club_hit(p):
    """Return (role, club, school) if the person leads an ETA/search-fund club at one of the six schools."""
    schools_attended = " ".join(((e.get("school") or {}).get("name") or "") for e in p.get("educations") or [])
    cands = [(t, c) for t, c, d in positions(p) if (d.get("end") is None or str(d.get("end", ""))[:4] >= "2025")]
    cands += [((o or {}).get("position") or "", (o or {}).get("name") or "") for o in p.get("organizations") or []]
    cands.append((cur_title(p), company_name(p)))
    for role, org in cands:
        blob = f"{role} {org}"
        if CLUB_RE.search(blob) and LEAD_RE.search(role):
            for school, rx in SCHOOLS:
                if re.search(rx, blob, re.I) or re.search(rx, schools_attended, re.I):
                    return role.strip(), org.strip(), school
    return None


# ------------------------------------------------------------------ hooks

def year(d):
    s = (d or {}).get("start") or ""
    return s[:4] if s else ""


def short(s, n=150):
    s = re.sub(r"\s+", " ", str(s or "")).strip()
    return s if len(s) <= n else s[: n - 1].rsplit(" ", 1)[0] + "…"


def year_note(ay):
    m = re.match(r"\d{4}-\d{2}", ay)
    if not m:
        return ""
    return f" ({m.group(0)} board" + ("; recent alum)" if "alum" in ay.lower() else ")")


def role_at(title, comp, since=""):
    title = re.sub(r"\s+", " ", title or "").strip(" |-")
    if comp and comp.lower() in title.lower():
        base = title
    else:
        base = f"{title} at {comp}" if comp else title
    return base + (f" (since {since})" if since else "")


def build_hook(seg, p):
    title, comp = cur_title(p), company_name(p)
    pos = positions(p)
    since = ""
    for t, c, d in pos:
        if c == comp and d.get("end") is None and year(d):
            since = year(d)
            break
    csum = ((p.get("company") or {}).get("summary") or {})
    desc = csum.get("description") if isinstance(csum, dict) else ""
    hq = ""
    loc = (p.get("company") or {}).get("location") or {}
    if isinstance(loc, dict):
        hq = (loc.get("hq") or loc).get("city", "") if isinstance(loc.get("hq") or loc, dict) else ""
    if seg.startswith("SBA"):
        return short(role_at(title, comp, since) + (f"; based in {hq}" if hq else ""))
    if seg.startswith("Search"):
        h = headline(p)
        m = re.search(r"[^.;|]*?(search fund|searcher|acqui[^.;|]*)[^.;|]*", h + " | " + str(desc or ""), re.I)
        fact = m.group(0).strip() if m else h
        return short(role_at(title, comp, since) + (f" — {fact}" if fact and fact.lower() not in (title + comp).lower() else ""))
    if seg.startswith("Small"):
        m = re.search(r"[^.;|]*(acqui|roll[- ]?up|hvac|plumb|pest|landscap|accounting|home service|cpa|bookkeep)[^.;|]*", str(desc or "") + " " + headline(p), re.I)
        fact = m.group(0).strip() if m else short(desc, 100)
        return short(role_at(title, comp, since) + (f" — {fact}" if fact else ""))
    hit = club_hit(p)
    if hit:
        role, club, school = hit
        return short(f"{role} of {club} at {school}")
    return short(role_at(title, comp))


# ------------------------------------------------------------------ main

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--per-segment", type=int, default=50)
    ap.add_argument("--max-pages", type=int, default=10, help="search pages (of 25) per query")
    ap.add_argument("--out", default=str(OUT))
    ap.add_argument("--segments", default="1,2,3,4", help="comma list of segment numbers to run")
    ap.add_argument("--append", action="store_true", help="keep rows already in --out and add to them")
    ap.add_argument("--cache", default=str(CACHE), help="response cache file (use one per parallel process)")
    ap.add_argument("--leaders-csv", default="", help="web-sourced club leaders (name,role,club,school,linkedin_url,...) to enrich for segment 4")
    args = ap.parse_args()
    if not API_KEY:
        sys.exit("no AI Ark key in env")

    names, links, emails, sources = load_already_contacted()
    print("Dedupe sources:\n  " + "\n  ".join(sources) + f"\n  => {len(names)} names, {len(links)} linkedin urls, {len(emails)} emails\n")

    api = AIArk(API_KEY, cache_file=Path(args.cache))
    rows, seen_ids, seen_names = [], set(), set()
    per_company = {}
    if args.append and Path(args.out).exists():
        with open(args.out, newline="", encoding="utf-8") as fh:
            rows = list(csv.DictReader(fh))
        for r in rows:
            seen_names.add(fold_name(r["Name"]))
            emails[r["Email"].lower()] = "this wave"
    stats = {}
    wanted = {int(x) for x in args.segments.split(",") if x.strip()}
    for si, seg in enumerate(SEGMENTS, 1):
        if si not in wanted:
            continue
        sname = seg["name"]
        st = stats[sname] = {"searched": 0, "filtered_out": 0, "dup_prev_wave": 0, "dup_in_run": 0, "exported": 0, "no_email": 0, "unverified_or_free": 0, "kept": 0, "dup_examples": []}
        for q in seg["queries"]:
            if st["kept"] >= args.per_segment:
                break
            q_kept, q_cap = 0, q.get("max", args.per_segment)
            for page in range(args.max_pages):
                if st["kept"] >= args.per_segment or q_kept >= q_cap:
                    break
                body = {**q["body"], "page": page, "size": PAGE_SIZE}
                status, data = api.post("/v1/people", body)
                if status != 200 or not isinstance(data, dict):
                    print(f"  [{sname}] {q['label']} p{page}: HTTP {status} {str(data)[:200]}")
                    break
                hits = data.get("content") or []
                print(f"  [{sname}] {q['label']} p{page}: {len(hits)} hits of {data.get('totalElements')}", flush=True)
                for p in hits:
                    st["searched"] += 1
                    pid = p.get("id")
                    full = (p.get("profile") or {}).get("full_name") or ""
                    li = clean_linkedin((p.get("link") or {}).get("linkedin") or "")
                    if (p.get("location") or {}).get("country") not in (None, "United States"):
                        st["filtered_out"] += 1
                        continue
                    if not seg["keep"](p):
                        st["filtered_out"] += 1
                        continue
                    fn = fold_name(full)
                    if fn in names or (li and li.lower() in links):
                        st["dup_prev_wave"] += 1
                        if len(st["dup_examples"]) < 5:
                            st["dup_examples"].append(f"{full} ({names.get(fn) or links.get(li.lower())})")
                        continue
                    if pid in seen_ids or fn in seen_names:
                        st["dup_in_run"] += 1
                        continue
                    ckey = fold_name(company_name(p))
                    if ckey and per_company.get((sname, ckey), 0) >= (20 if "live oak" in ckey else 2):
                        st["filtered_out"] += 1
                        continue
                    seen_ids.add(pid)
                    st["exported"] += 1
                    es, ed = api.export_by_id(pid)
                    person = unwrap(ed) if es == 200 else None
                    if not person:
                        st["no_email"] += 1
                        continue
                    email, conf = pick_email(person)
                    outputs = (person.get("email") or {}).get("output") if isinstance(person.get("email"), dict) else person.get("email")
                    best = next((o for o in (outputs or []) if o.get("address") == email), {}) if isinstance(outputs, list) else {}
                    if not email or conf != "high" or best.get("free") or email.lower() in emails:
                        st["unverified_or_free"] += 1 if email else 0
                        st["no_email"] += 0 if email else 1
                        continue
                    seen_names.add(fn)
                    per_company[(sname, ckey)] = per_company.get((sname, ckey), 0) + 1
                    st["kept"] += 1
                    q_kept += 1
                    rows.append({
                        "Name": full,
                        "Title": cur_title(p),
                        "Company": company_name(p),
                        "LinkedIn URL": li or clean_linkedin((person.get("link") or {}).get("linkedin") or ""),
                        "Source Variant": f"W3/{q['label']}",
                        "Email": email,
                        "Email Confidence": conf,
                        "Phone": "",
                        "Enrichment Status": "email_only (aiark_search)",
                        "Segment": sname,
                        "Hook": build_hook(sname, p),
                    })
                    if st["kept"] >= args.per_segment or q_kept >= q_cap:
                        break
                if data.get("last") or len(hits) < PAGE_SIZE:
                    break

    if args.leaders_csv and 4 in wanted:
        sname = "MBA ETA club leaders"
        st = stats.setdefault(sname, {"searched": 0, "filtered_out": 0, "dup_prev_wave": 0, "dup_in_run": 0, "exported": 0, "no_email": 0, "unverified_or_free": 0, "kept": 0, "dup_examples": []})
        with open(args.leaders_csv, newline="", encoding="utf-8") as fh:
            leaders = list(csv.DictReader(fh))
        for L in leaders:
            if st["kept"] >= args.per_segment:
                break
            name, school = (L.get("name") or "").strip(), (L.get("school") or "").strip()
            if not name:
                continue
            st["searched"] += 1
            fn = fold_name(name)
            if fn in names:
                st["dup_prev_wave"] += 1
                continue
            if fn in seen_names:
                st["dup_in_run"] += 1
                continue
            person, li = None, clean_linkedin(L.get("linkedin_url") or "")
            if li:
                es, ed = api.export_by_linkedin(li)
                person = unwrap(ed) if es == 200 else None
            if not person:
                body = {"contact": {"fullName": {"mode": "SMART", "include": [name]}}, "page": 0, "size": 5}
                ss, sd = api.post("/v1/people", body)
                hits = (sd or {}).get("content") or [] if ss == 200 and isinstance(sd, dict) else []
                last = name.split()[-1].lower()
                for h in hits:
                    hp = h.get("profile") or {}
                    if last not in (hp.get("full_name") or "").lower():
                        continue
                    blob = person_text(h)
                    if not re.search(dict(SCHOOLS)[school] if school in dict(SCHOOLS) else re.escape(school), blob, re.I):
                        continue
                    st["exported"] += 1
                    es, ed = api.export_by_id(h.get("id"))
                    person = unwrap(ed) if es == 200 else None
                    if person:
                        li = li or clean_linkedin((h.get("link") or {}).get("linkedin") or "")
                        break
            if not person:
                st["no_email"] += 1
                continue
            email, conf = pick_email(person)
            outputs = (person.get("email") or {}).get("output") if isinstance(person.get("email"), dict) else person.get("email")
            best = next((o for o in (outputs or []) if o.get("address") == email), {}) if isinstance(outputs, list) else {}
            if not email or conf != "high" or best.get("free") or email.lower() in emails:
                st["unverified_or_free"] += 1 if email else 0
                st["no_email"] += 0 if email else 1
                continue
            seen_names.add(fn)
            emails[email.lower()] = "this wave"
            st["kept"] += 1
            pt = (person.get("profile") or {})
            rows.append({
                "Name": pt.get("full_name") or name,
                "Title": f"{L.get('role', '').strip()}, {L.get('club', '').strip()}".strip(", ") or cur_title(person),
                "Company": school,
                "LinkedIn URL": li or clean_linkedin((person.get("link") or {}).get("linkedin") or ""),
                "Source Variant": "W3/eta-club-web",
                "Email": email,
                "Email Confidence": conf,
                "Phone": "",
                "Enrichment Status": "email_only (aiark_lookup)",
                "Segment": sname,
                "Hook": short(f"{L.get('role', '').strip()} of {L.get('club', '').strip()} at {school}" + year_note(L.get("academic_year") or "")),
            })

    with open(args.out, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=FIELDS)
        w.writeheader()
        w.writerows(rows)

    print(f"\nWrote {args.out}: {len(rows)} rows. API calls this run: {api.calls}\n")
    print(f"{'Segment':34} {'kept':>5} {'searched':>8} {'filtered':>8} {'dup_prev':>8} {'dup_run':>7} {'exported':>8} {'no_email':>8} {'unverif':>7}")
    for s, st in stats.items():
        print(f"{s:34} {st['kept']:>5} {st['searched']:>8} {st['filtered_out']:>8} {st['dup_prev_wave']:>8} {st['dup_in_run']:>7} {st['exported']:>8} {st['no_email']:>8} {st['unverified_or_free']:>7}")
    for s, st in stats.items():
        if st["dup_examples"]:
            print(f"  {s} dup examples: " + "; ".join(st["dup_examples"]))
    Path(args.out).with_suffix(".stats.json").write_text(json.dumps(stats, indent=1))


if __name__ == "__main__":
    main()
